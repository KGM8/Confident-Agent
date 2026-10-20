"""
control_support_check.py -- does each positive control's excerpt actually SUPPORT the answer?

A positive control in tools/critic_controls.py only guarantees that the gold answer STRING appears
in a retrieved chunk of the question's source document. That is necessary but not sufficient for
support: "1976" or "ANC" can appear in a sentence about something else. This tool lists, for every
positive control, the retrieved sentence(s) containing the gold string, and writes a small sheet
for a person to judge whether that sentence supports the answer to the question. No LLM calls.

    python tools/control_support_check.py
    -> results/control_support_check.xlsx   (fill column F: yes / no / unclear)
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

import critic_audit as ca  # noqa: E402


def gold_sentences(chunks, gold, expected):
    g = ca.norm(gold)
    out = []
    for c in chunks:
        if c.chunk.source_id != expected:
            continue
        for s in re.split(r"(?<=[.!?])\s+|\n+", c.chunk.text):
            if g and g in ca.norm(s):
                out.append(s.strip()[:400])
    return out


def main():
    from main import build_agent, load_config
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    agent = build_agent(load_config("config/config.yaml"))
    controls = [r for r in ca.load_jsonl("results/critic_controls.jsonl") if r["type"] == "positive"]
    samples = {}
    for kind, path in (("mcq", "mcq_sample30.jsonl"), ("open", "open_sample30.jsonl")):
        for r in ca.load_jsonl(path):
            samples[(kind, r["id"])] = r
    wb = Workbook(); ws = wb.active; ws.title = "Support check"
    ws.append(["Benchmark", "Id", "Question", "Answer judged by the critic", "Retrieved sentence(s) containing the gold string",
               "Does the sentence support this answer to this question? (yes / no / unclear)", "Critic verdict"])
    for c in controls:
        s = samples[(c["kind"], c["id"])]
        gold = s["choices"][s["answer"]] if c["kind"] == "mcq" else s["answer"]
        expected = s["source_document"].replace(".md", "")
        sents = gold_sentences(agent.retriever.retrieve(c["question"]).chunks, gold, expected)
        ws.append([c["kind"], c["id"], c["question"], c["answer"], "\n---\n".join(sents) or "(none found)", None, c["robust_verdict"]])
    for col, w in zip("ABCDEFG", (10, 6, 45, 30, 70, 22, 14)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6).fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"F2:F{ws.max_row}")
    ws.freeze_panes = "C2"
    os.makedirs("results", exist_ok=True)
    wb.save("results/control_support_check.xlsx")
    print(f"wrote results/control_support_check.xlsx ({len(controls)} positive controls). Fill column F, then send it back.")


if __name__ == "__main__":
    main()
